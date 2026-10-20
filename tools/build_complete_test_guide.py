"""Build the Version 2 test guide from the current Studio workflow."""

from datetime import date
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT, WD_TAB_LEADER
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / "docs/guides/Delphi App Translation Studio Complete Test Guide.docx"


def add_toc(document):
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.tab_stops.add_tab_stop(
        Inches(6.5), WD_TAB_ALIGNMENT.RIGHT, WD_TAB_LEADER.DOTS
    )
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instruction = OxmlElement("w:instrText")
    instruction.set(qn("xml:space"), "preserve")
    instruction.text = 'TOC \\o "1-2" \\h \\z \\u'
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    run._r.extend([begin, instruction, separate])
    paragraph.add_run(
        "Test environment and evidence\t2\n"
        "Project opening and scan\t2\n"
        "Source and target language pairs\t2\n"
        "Adding and reviewing several languages\t3\n"
        "Glossary and translation review\t3\n"
        "Validation export and integration\t3\n"
        "Running the translated application\t4\n"
        "Automated regression and final decision\t4"
    )
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    paragraph.add_run()._r.append(end)


def step(document, action, expected):
    paragraph = document.add_paragraph(style="List Number")
    paragraph.add_run(action + " ")
    paragraph.add_run("Expected: " + expected).italic = True


def check(document, text):
    document.add_paragraph(text, style="List Bullet")


def add_section(document, title, introduction):
    document.add_heading(title, level=1)
    document.add_paragraph(introduction)


def make_document():
    doc = Document()
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.85)
    section.bottom_margin = Inches(0.75)
    section.left_margin = Inches(0.9)
    section.right_margin = Inches(0.9)

    normal = doc.styles["Normal"]
    normal.font.name = "Arial"
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = RGBColor(0, 0, 0)
    normal.paragraph_format.space_after = Pt(7)
    for style_name, size in (("Title", 23), ("Heading 1", 15), ("Heading 2", 12)):
        style = doc.styles[style_name]
        style.font.name = "Arial"
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.font.bold = style_name != "Title"
        style.paragraph_format.space_before = Pt(15 if style_name != "Title" else 0)
        style.paragraph_format.space_after = Pt(7)
        style.paragraph_format.keep_with_next = True
        if style_name == "Title":
            ppr = style.element.get_or_add_pPr()
            borders = ppr.find(qn("w:pBdr"))
            if borders is not None:
                ppr.remove(borders)

    title = doc.add_paragraph(style="Title")
    title.add_run("Delphi App Translation Studio Complete Test Guide")
    subtitle = doc.add_paragraph()
    subtitle.add_run("Version 2.0  |  Windows VCL and FireMonkey").bold = True
    doc.add_paragraph("Last changed: " + date.today().strftime("%B %d, %Y"))
    doc.add_paragraph(
        "This guide gives a repeatable acceptance test for the Version 2 Studio, "
        "its generated language packs, and a Delphi application that consumes them. "
        "Run the tests against a disposable, version-controlled copy of the target "
        "application. A successful compiler build alone does not establish that a "
        "translation is correct or that every screen remains usable."
    )
    doc.add_paragraph(
        "Record the result of each section. Stop before distribution if a target "
        "file changes unexpectedly, a language pack is incompatible, validation "
        "reports errors, or a translated screen fails review."
    )
    doc.add_heading("Contents", level=1)
    add_toc(doc)
    doc.add_page_break()

    add_section(doc, "Test environment and evidence",
        "Test the exact RAD Studio version, framework, platform, and build "
        "configuration intended for the application. Keep Studio, packages, "
        "and generated dependencies from one Version 2 build together.")
    for item in (
        "Record the Studio build label, RAD Studio version, VCL or FMX, Win32 or Win64, and Debug or Release.",
        "Record the target project path and Git revision. Save hashes of its DPR, DPROJ, PAS, DFM, and FMX files before testing.",
        "Use test text only. Do not place API keys, customer data, or private source in screenshots or shared logs.",
        "Use a provider account and key supplied by the project owner. Confirm quota and possible charges before a multi-language request.",
    ):
        check(doc, item)
    doc.add_heading("Results to retain", level=2)
    doc.add_paragraph(
        "Keep the Studio validation output, per-language catalog and pack paths, "
        "target build logs, source-file hash comparison, selected screenshots, "
        "and an issue list. Record failures with the exact action, expected result, "
        "actual result, and affected language pair."
    )

    add_section(doc, "Project opening and scan",
        "The Original app language is the language already present in the saved "
        "forms and resources. It is not the language you want to add.")
    step(doc, "Open a saved Delphi VCL or FMX project in the Studio.",
        "The Project page shows the correct project and advances to Scan; Scan Project is available after an original language is chosen.")
    step(doc, "Select the Original app language and scan.",
        "The scan lists source items and counts; no target project source file is rewritten.")
    step(doc, "Edit a saved form or source file after scanning, then attempt translation.",
        "The Studio requires a fresh scan instead of using stale text. Restore the test file and scan again.")
    step(doc, "Compare target project source-file hashes with the starting values.",
        "They match unless the tester intentionally edited the target project.")

    add_section(doc, "Source and target language pairs",
        "Use at least one non-English authored project. GermanProject is a useful "
        "VCL fixture, but it does not replace testing the developer's real application.")
    step(doc, "Translate a German-authored project to a different supported language, such as Dutch.",
        "The catalog identifies de-DE as source and the selected target locale as target. German source text is preserved.")
    step(doc, "Choose an identical source and target locale.",
        "The Studio rejects the pair before a provider request.")
    step(doc, "Try a same-base regional pair, such as fr-FR to fr-CA, with Google Basic.",
        "The Studio explains that automatic translation is unavailable for the provider-mapped pair; it does not claim a completed translation.")
    step(doc, "Confirm Traditional Chinese is offered as a source and Simplified Chinese is target-only.",
        "The source and target lists reflect the documented Version 2 scope.")
    step(doc, "Create the same target from two different authored source languages in separate test projects.",
        "Catalogs, glossaries, and packs retain separate source-target identities; one pair never reuses another pair's reviewed wording silently.")

    add_section(doc, "Adding and reviewing several languages",
        "The Add languages page creates draft catalogs; it does not deploy packs. "
        "Provider requests can incur charges, so start with a small set.")
    step(doc, "After scanning, open Add languages from the orange-accented left panel and select two or three targets.",
        "The source is visible and the target count matches the selection.")
    step(doc, "With Google and DeepL in turn, use Select All from an English source.",
        "The source and same-base English regional target are excluded. Clear resets the selection.")
    step(doc, "Translate the small set and watch the progress display.",
        "The display names the language being processed; each successful catalog is saved as a draft.")
    step(doc, "Use Next: Review, then Save and next for each language.",
        "The editor shows the current target and its place in the batch. A saved correction persists when the catalog is reopened.")
    step(doc, "Cancel or simulate a later-language provider failure after one draft is saved.",
        "The saved language remains accessible through Next: Review. The failed or unprocessed language is not presented as complete.")
    step(doc, "Finish without exporting, then run the target application.",
        "Its language selector is unchanged because drafts were not deployed.")

    add_section(doc, "Glossary and translation review",
        "A glossary belongs to one project and one source-target pair. Machine "
        "output remains a draft until a person checks its meaning and layout.")
    step(doc, "Open Glossary for a non-English project before opening a translation catalog.",
        "The glossary uses the Original app language. If that language is not selected, the Studio asks for it instead of assuming English.")
    step(doc, "Save a preferred term and reopen the matching glossary.",
        "The term persists only for that project and language pair.")
    step(doc, "Inspect representative short words, placeholders, accelerator keys, dates, currency, and long labels in each generated catalog.",
        "The review states are deliberate; a native or fluent reviewer confirms meaning before approval.")
    step(doc, "Review Arabic, Hebrew, or another RTL target when in scope.",
        "Text direction, punctuation, mixed Latin text, wrapping, and control order are readable in the running application.")

    add_section(doc, "Validation export and integration",
        "Validation checks catalog structure. Warnings need review; errors must "
        "be corrected before export. Export is a separate choice.")
    step(doc, "Run Validate All Languages after reviewing the batch.",
        "Each target has a separate error and warning count; the action to correct errors is clear.")
    step(doc, "Introduce a missing required translation or broken placeholder in a disposable catalog and revalidate.",
        "The error blocks export. Restore the entry and validate again.")
    step(doc, "Choose Done for now after validation.",
        "Drafts remain saved and no runtime pack is installed.")
    step(doc, "Choose Export and deploy languages after validation passes.",
        "A please-wait display identifies the current export language and integration stage; each pack has the expected source and target metadata.")
    step(doc, "Inspect the generated component kit and target output folders.",
        "Only compatible packs from the chosen authored source are deployed. An incompatible older pack is rejected, not silently mixed or deleted.")
    step(doc, "Compare target source-file hashes again and inspect the DPROJ Search Path.",
        "The Studio did not rewrite source or project files. Any permanent Search Path change remains a developer-owned IDE setting.")

    add_section(doc, "Running the translated application",
        "A runtime test is required for every framework and platform the project "
        "will ship. Repeat these checks on representative complex forms.")
    for item in (
        "Build and launch the target in each intended Win32 or Win64 Debug or Release configuration.",
        "Confirm the selector lists the authored source plus only deployed, compatible targets. A draft catalog alone must not add a language.",
        "Switch between source and target repeatedly; verify labels, menus, dialogs, memo text, status text, and dynamic text.",
        "Restart with a saved target-language preference. Confirm the selected language and visible text return correctly.",
        "Check long text, controls near window edges, high DPI, resizing, dates, numbers, and RTL layouts where applicable.",
        "Ask a fluent speaker to review meaning, tone, grammar, and terminology for each language intended for distribution.",
    ):
        check(doc, item)

    add_section(doc, "Automated regression and final decision",
        "Run the repository's complete local validation suite after source, form, "
        "package, or guide changes. Preserve its output and do not equate a "
        "passing automated test with linguistic approval.")
    doc.add_paragraph(
        "Run tools\\tests\\RunPhase10ReleaseValidation.ps1 from the Version 2 "
        "checkout with the RAD Studio 13 Florence environment available. "
        "Also run tools\\check_artifacts_current.ps1 and "
        "tools\\build_source_distribution.ps1 -ValidateOnly. The source "
        "distribution check must not create or publish an archive."
    )
    doc.add_heading("Acceptance record", level=2)
    for item in (
        "Studio build and Git revision: ______________________________",
        "Target project and source language: _________________________",
        "Targets, provider, framework, and platforms: _______________",
        "Automated suite result and log location: ____________________",
        "Native-language and UI review result: ______________________",
        "Open issues and release decision: ___________________________",
    ):
        doc.add_paragraph(item)

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    footer.add_run("Delphi App Translation Studio  |  Version 2.0")
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    make_document()
